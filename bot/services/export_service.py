"""Экспорт таблиц БД в .xlsx — поверх SQLite, не меняя основную логику.

Данные читаются через ORM-модели, поэтому типы (datetime/int/str)
сохраняются корректно автоматически. openpyxl пишет UTF-8 нативно.
"""
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select

from bot.database.engine import session_maker
from bot.database.models import Auction, Bid, LootHistory, Player, Transaction

# Папка для выгрузок
EXPORT_DIR = Path("exports")

# Реестр таблиц: техническое имя -> (ORM-модель, человекочитаемое имя листа)
TABLE_REGISTRY: dict[str, tuple[type, str]] = {
    "players": (Player, "Игроки"),
    "transactions": (Transaction, "Операции"),
    "auctions": (Auction, "Аукционы"),
    "bids": (Bid, "Ставки"),
    "loot_history": (LootHistory, "История лута"),
}

# Русские заголовки колонок (если колонки нет — берётся техническое имя)
COLUMN_LABELS: dict[str, str] = {
    "id": "ID",
    "discord_id": "Discord ID",
    "player_id": "ID игрока",
    "auction_id": "ID аукциона",
    "winner_id": "ID победителя",
    "officer_id": "Discord ID офицера",
    "started_by": "Запустил (Discord ID)",
    "channel_id": "ID канала",
    "message_id": "ID сообщения",
    "display_name": "Игрок",
    "current_dkp": "Текущий DKP",
    "total_earned": "Всего заработано",
    "total_spent": "Всего потрачено",
    "amount": "Сумма",
    "balance_after": "Баланс после",
    "type": "Тип",
    "reason": "Причина",
    "item_name": "Предмет",
    "min_bid": "Мин. ставка",
    "bid_step": "Шаг ставки",
    "image_url": "Изображение",
    "status": "Статус",
    "winning_bid": "Выигрышная ставка",
    "cost": "Стоимость",
    "created_at": "Создано",
    "finished_at": "Завершено",
}

# Человекочитаемые значения для enum-полей
TYPE_LABELS = {
    "EARN": "Начисление",
    "SPEND": "Списание",
    "PENALTY": "Штраф",
    "AUCTION": "Аукцион",
    "DECAY": "Снижение",
    "ACTIVE": "Активен",
    "FINISHED": "Завершён",
    "CANCELLED": "Отменён",
}

# Excel хранит числа как double и теряет точность выше 2^53.
# Discord ID (~10^18) больше этого предела — пишем такие как текст.
_MAX_SAFE_INT = 2**53

# Стили оформления
_HEADER_FILL = PatternFill("solid", fgColor="2F5597")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center")
_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_DATE_FORMAT = "DD.MM.YYYY HH:MM"


def _label_for_column(col_name: str) -> str:
    return COLUMN_LABELS.get(col_name, col_name)


def _humanize_value(col_name: str, value):
    """Делает значения дружелюбнее для человека."""
    if value is None:
        return ""
    if col_name in ("type", "status") and isinstance(value, str):
        return TYPE_LABELS.get(value, value)
    # Большие ID (Discord snowflake) — строкой, иначе Excel теряет точность
    if isinstance(value, int) and abs(value) > _MAX_SAFE_INT:
        return str(value)
    return value


async def _fetch_table(model: type) -> tuple[list[str], list[list]]:
    """Возвращает (имена колонок, строки данных) через ORM — типы сохраняются."""
    columns = [c.name for c in model.__table__.columns]
    async with session_maker() as session:
        result = await session.execute(select(model))
        rows = result.scalars().all()

    data: list[list] = []
    for row in rows:
        data.append([_humanize_value(c, getattr(row, c)) for c in columns])
    return columns, data


def _write_sheet(ws: Worksheet, columns: list[str], rows: list[list]) -> None:
    # Шапка
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=_label_for_column(col_name))
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGN
        cell.border = _BORDER

    # Данные
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.border = _BORDER
            if isinstance(value, (datetime, date)):
                cell.number_format = _DATE_FORMAT
                cell.alignment = Alignment(horizontal="center")

    # Автоширина колонок (с разумным потолком)
    for c_idx, col_name in enumerate(columns, start=1):
        max_len = len(_label_for_column(col_name))
        for row in rows:
            val = row[c_idx - 1]
            if isinstance(val, (datetime, date)):
                length = 16
            else:
                length = len(str(val))
            max_len = max(max_len, length)
        ws.column_dimensions[get_column_letter(c_idx)].width = min(max_len + 3, 50)

    # Заморозка шапки + автофильтр
    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(columns))
    last_row = len(rows) + 1
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"


async def export_to_excel(
    table_names: list[str] | None = None,
    path: str | Path | None = None,
) -> Path:
    """Экспортирует таблицы в один .xlsx (по листу на таблицу).

    :param table_names: список технических имён таблиц; None — все таблицы.
    :param path: путь к файлу; None — авто-имя в папке exports/.
    :return: путь к созданному файлу.
    """
    if table_names is None:
        table_names = list(TABLE_REGISTRY.keys())

    # Валидация
    unknown = [t for t in table_names if t not in TABLE_REGISTRY]
    if unknown:
        raise ValueError(f"Неизвестные таблицы: {', '.join(unknown)}")

    if path is None:
        EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        suffix = "all" if len(table_names) == len(TABLE_REGISTRY) else "_".join(table_names)
        path = EXPORT_DIR / f"dkp_export_{suffix}_{stamp}.xlsx"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)  # убираем дефолтный пустой лист

    for table_name in table_names:
        model, sheet_title = TABLE_REGISTRY[table_name]
        columns, rows = await _fetch_table(model)
        ws = wb.create_sheet(title=sheet_title)
        if not rows:
            # Всё равно пишем шапку, чтобы лист был осмысленным
            _write_sheet(ws, columns, [])
        else:
            _write_sheet(ws, columns, rows)

    wb.save(path)
    return path
